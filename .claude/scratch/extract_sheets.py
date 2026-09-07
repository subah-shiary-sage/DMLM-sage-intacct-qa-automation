import openpyxl
import json
import os

SRC = r"Test_Cases_and_Checklists\SNL Test cases.xlsx"
OUT_DIR = r".claude\scratch\snl_audit"

wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)

os.makedirs(OUT_DIR, exist_ok=True)

manifest = []

for sheet_name in wb.sheetnames:
    if sheet_name == "Summary":
        continue
    ws = wb[sheet_name]
    rows = []
    for r in range(1, ws.max_row + 1):
        row_vals = []
        has_content = False
        for c in range(1, 11):  # A-J
            v = ws.cell(row=r, column=c).value
            row_vals.append(v)
            if v is not None and str(v).strip() != "":
                has_content = True
        if has_content:
            rows.append({"row": r, "values": row_vals})

    safe_name = sheet_name.replace(" ", "_").replace("/", "-")
    out_path = os.path.join(OUT_DIR, f"{safe_name}.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump({"sheet": sheet_name, "rows": rows}, f, indent=1, default=str)

    manifest.append({"sheet": sheet_name, "file": f"{safe_name}.json", "row_count": len(rows)})

with open(os.path.join(OUT_DIR, "_manifest.json"), "w", encoding="utf-8") as f:
    json.dump(manifest, f, indent=1)

print(f"Extracted {len(manifest)} sheets")
for m in manifest:
    print(m["sheet"], "->", m["row_count"], "content rows")
