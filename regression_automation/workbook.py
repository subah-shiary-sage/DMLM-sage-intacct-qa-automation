"""
workbook.py — Read/write access to Regression/LME_Regression_Consolidated.xlsx.

Every layout quirk of that workbook is encoded here once, so callers never
hardcode row/column numbers:

  * Two checklist sheet families with different header rows and columns.
    ``Loan Interest Rate - CRUD`` is an exception within the CRUD family
    (header on row 1, not 6).
  * Section-separator rows ("CREATE", "EDIT", "Page load & header", ...) recur
    *throughout* each sheet, not just after the header. A row is real test data
    only when its objective cell is non-empty.
  * ``Loan Type - CRUD`` ends with a FINDINGS / EDGE-CASE ADDENDUM narrative
    block whose rows look like data but aren't re-runnable scenarios. Excluded
    per-sheet, never globally: other sheets legitimately use non-numeric SL#s
    (e.g. "LT-LISTER-DUPCOL" in Config Listers is a real Passed row).
  * Stored counts go stale. Per-sheet header blocks and the Summary table are
    always recomputed from the live status column, never read and trusted.
"""

from __future__ import annotations

import datetime as _dt
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator, Optional

from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.workbook import Workbook

WORKBOOK_PATH = Path("Regression/LME_Regression_Consolidated.xlsx")

# The four values the status dropdowns are constrained to.
STATUS_PASSED = "Passed"
STATUS_FAILED = "Failed"
STATUS_UNTESTED = "Untested"
STATUS_TO_BE_TESTED = "To be tested"
VALID_STATUSES = (STATUS_PASSED, STATUS_FAILED, STATUS_UNTESTED, STATUS_TO_BE_TESTED)

# Statuses the runner treats as "not yet executed".
PENDING_STATUSES = (STATUS_UNTESTED, STATUS_TO_BE_TESTED)

STATUS_FILL = {
    STATUS_PASSED: "C6EFCE",
    STATUS_FAILED: "FFC7CE",
    STATUS_UNTESTED: "D9D9D9",
    STATUS_TO_BE_TESTED: "FFEB9C",
}

JIRA_BUGS_SHEET = "Jira Bugs (IADSSL-1714)"
JIRA_BUGS_HEADER_ROW = 2
JIRA_BUGS_FIRST_DATA_ROW = 3
BUG_PENDING_FILL = "FFEB9C"  # yellow: found but not yet filed to JIRA

SUMMARY_SHEET = "Summary"
SUMMARY_FIRST_ROW = 9
SUMMARY_TOTAL_ROW = 20

_CRUD = dict(header_row=6, status_col=6, comment_col=7, objective_col=2)
_CRUD_H1 = dict(header_row=1, status_col=6, comment_col=7, objective_col=2)
_LISTER = dict(header_row=5, status_col=3, comment_col=4, objective_col=2)

# Sheet name -> layout. Order matches the workbook's tab order, which is also the
# order the Summary coverage table (rows 9-19) expects.
FAMILY_LAYOUT: dict[str, dict] = {
    "Loan Account - CRUD": {**_CRUD},
    "Loan Account - Lister": {**_LISTER},
    "Loan Type - CRUD": {
        **_CRUD,
        # FINDINGS / EDGE-CASE ADDENDUM block at the end of the sheet: narrative
        # notes, not independently re-runnable scenarios.
        "exclude_col1_pattern": r"^(F\d+|E\d+|NOTE)$",
    },
    "Loan Interest Rate - CRUD": {**_CRUD_H1},
    "Loan Category - CRUD": {**_CRUD},
    "Loan Fee Type - CRUD": {**_CRUD},
    "Loan Adjustment - CRUD": {**_CRUD},
    "Config Listers LT-LIR-LC-LFT": {**_LISTER},
    "Edit-View-Delete + Update": {**_LISTER},
    "Loan Statements - Lister": {**_LISTER},
    "Lending Workbench - Lister": {**_LISTER},
}

# Sheets in Summary-table order (rows 9-19).
SUMMARY_SHEET_ORDER = list(FAMILY_LAYOUT)


class WorkbookLockedError(RuntimeError):
    """The workbook is open in Excel, so it cannot be written."""


@dataclass
class TestRow:
    sheet: str
    row: int
    sl_no: str | int | None
    objective: str
    status: Optional[str]
    comment: Optional[str]

    @property
    def is_pending(self) -> bool:
        return (self.status or "").strip() in PENDING_STATUSES


def check_not_locked(path: Path = WORKBOOK_PATH) -> None:
    """
    Raise before doing any work if Excel holds the workbook open.

    Excel writes a sibling ``~$<name>`` lock file. Catching it up front turns a
    late, confusing PermissionError (after the browser work and possibly after
    JIRA posts) into an actionable message before anything runs.
    """
    path = Path(path)
    lock = path.parent / f"~${path.name}"
    if lock.exists():
        raise WorkbookLockedError(
            f"{path.name} appears to be open in Excel (lock file {lock.name} present).\n"
            "Close the workbook and re-run."
        )


def _layout(sheet_name: str) -> dict:
    try:
        return FAMILY_LAYOUT[sheet_name]
    except KeyError:
        raise KeyError(
            f"Unknown checklist sheet {sheet_name!r}. "
            f"Known sheets: {', '.join(FAMILY_LAYOUT)}"
        ) from None


def iter_test_rows(
    wb: Workbook,
    sheet_name: str,
    row_start: int | None = None,
    row_end: int | None = None,
) -> Iterator[TestRow]:
    """
    Yield the real test-case rows on a sheet.

    Rows whose objective cell is empty are section separators and are skipped.
    ``row_start``/``row_end`` scope the walk to one module's block on a sheet
    shared by several modules (Config Listers).
    """
    layout = _layout(sheet_name)
    ws = wb[sheet_name]
    exclude = layout.get("exclude_col1_pattern")
    exclude_re = re.compile(exclude) if exclude else None

    first = row_start if row_start is not None else layout["header_row"] + 1
    last = row_end if row_end is not None else ws.max_row

    for r in range(first, last + 1):
        objective = ws.cell(r, layout["objective_col"]).value
        if objective is None or not str(objective).strip():
            continue  # section separator / blank
        sl = ws.cell(r, 1).value
        if exclude_re and sl is not None and exclude_re.match(str(sl).strip()):
            continue
        status = ws.cell(r, layout["status_col"]).value
        comment = ws.cell(r, layout["comment_col"]).value
        yield TestRow(
            sheet=sheet_name,
            row=r,
            sl_no=sl,
            objective=str(objective).strip(),
            status=str(status).strip() if status is not None else None,
            comment=str(comment).strip() if comment is not None else None,
        )


def pending_rows(
    wb: Workbook,
    sheet_name: str,
    *,
    include_failed: bool = False,
    row_start: int | None = None,
    row_end: int | None = None,
) -> list[TestRow]:
    """Rows still to execute: Untested / To be tested, plus Failed if asked."""
    wanted = set(PENDING_STATUSES)
    if include_failed:
        wanted.add(STATUS_FAILED)
    return [
        row
        for row in iter_test_rows(wb, sheet_name, row_start, row_end)
        if (row.status or "").strip() in wanted
    ]


def evidence_note(
    verdict: str,
    *,
    environment: str,
    run_date: str | None = None,
    existing: str | None = None,
    retest_ticket: str | None = None,
) -> str:
    """
    Build an evidence note in the convention already used throughout the workbook.

    New note:  ``EXECUTED <date> (<environment>): <verdict>``
    Retest:    appended to the existing cell as ``  ||  RETEST (<ticket>, <date>): <verdict>``
               so prior evidence is never overwritten.
    """
    date = run_date or _dt.date.today().isoformat()
    if existing and existing.strip():
        ticket = f"{retest_ticket}, " if retest_ticket else ""
        return f"{existing.strip()}  ||  RETEST ({ticket}{date}): {verdict}"
    return f"EXECUTED {date} ({environment}): {verdict}"


def write_result(wb: Workbook, row: TestRow, *, status: str, comment: str) -> None:
    """Write a row's status (with its fill colour) and evidence comment."""
    if status not in VALID_STATUSES:
        raise ValueError(
            f"Invalid status {status!r}; must be one of {', '.join(VALID_STATUSES)}"
        )
    layout = _layout(row.sheet)
    ws = wb[row.sheet]
    ws.cell(row.row, layout["status_col"]).value = status
    ws.cell(row.row, layout["status_col"]).fill = PatternFill(
        "solid", fgColor=STATUS_FILL[status]
    )
    ws.cell(row.row, layout["comment_col"]).value = comment


def count_statuses(wb: Workbook, sheet_name: str) -> dict[str, int]:
    """Live tally of a sheet's status column. Never reads a cached count."""
    counts = {s: 0 for s in VALID_STATUSES}
    counts["(blank)"] = 0
    for row in iter_test_rows(wb, sheet_name):
        key = (row.status or "").strip()
        counts[key if key in counts else "(blank)"] += 1
    return counts


def recompute_summary(wb: Workbook) -> None:
    """
    Rewrite the Summary coverage table (rows 9-19) and its TOTAL row from live
    per-sheet tallies.

    Only the numeric table is touched. The free-text DONE / NOT DONE / NOTES
    bullets below it are hand-maintained and left alone.
    """
    summary = wb[SUMMARY_SHEET]
    totals = [0, 0, 0, 0]

    for offset, sheet_name in enumerate(SUMMARY_SHEET_ORDER):
        counts = count_statuses(wb, sheet_name)
        items = sum(counts[s] for s in VALID_STATUSES) + counts["(blank)"]
        passed = counts[STATUS_PASSED]
        failed = counts[STATUS_FAILED]
        untested = counts[STATUS_UNTESTED] + counts[STATUS_TO_BE_TESTED]

        r = SUMMARY_FIRST_ROW + offset
        summary.cell(r, 2).value = items
        summary.cell(r, 3).value = passed
        summary.cell(r, 4).value = failed
        summary.cell(r, 5).value = untested

        totals[0] += items
        totals[1] += passed
        totals[2] += failed
        totals[3] += untested

    for col, value in zip((2, 3, 4, 5), totals):
        summary.cell(SUMMARY_TOTAL_ROW, col).value = value


def next_bug_number(wb: Workbook) -> int:
    """
    Next free bug number per the workbook's Jira Bugs sheet.

    The Bug # column is not reliably an int — provisional rows hold strings like
    "89 (provisional)" — so values are coerced and non-numeric ones ignored
    rather than allowed to break ``max()``.
    """
    ws = wb[JIRA_BUGS_SHEET]
    numbers: list[int] = []
    for r in range(JIRA_BUGS_FIRST_DATA_ROW, ws.max_row + 1):
        raw = ws.cell(r, 1).value
        if raw is None:
            continue
        if isinstance(raw, int):
            numbers.append(raw)
            continue
        match = re.match(r"\s*(\d+)", str(raw))
        if match:
            numbers.append(int(match.group(1)))
    return (max(numbers) + 1) if numbers else 1


def append_bug_row(
    wb: Workbook,
    *,
    bug_no: int | str,
    module_area: str,
    summary: str,
    status: str = "Open",
    related_checklist: str = "",
    pending: bool = False,
) -> int:
    """
    Mirror a defect into the Jira Bugs sheet. Returns the row it was written to.

    ``pending=True`` marks a defect that has been found but not yet filed to
    JIRA — shaded yellow so it is visible at a glance in the workbook, with the
    blocker recorded in the Status cell rather than the row being omitted.
    """
    ws = wb[JIRA_BUGS_SHEET]
    r = ws.max_row + 1
    values = (bug_no, module_area, summary, status, related_checklist)
    for col, value in enumerate(values, start=1):
        ws.cell(r, col).value = value
        if pending:
            ws.cell(r, col).fill = PatternFill("solid", fgColor=BUG_PENDING_FILL)
    if pending:
        ws.cell(r, 4).font = Font(bold=True)
    return r


def find_bug_row(wb: Workbook, bug_no: int) -> int | None:
    """Row holding this bug number, if already mirrored (keeps writes idempotent)."""
    ws = wb[JIRA_BUGS_SHEET]
    for r in range(JIRA_BUGS_FIRST_DATA_ROW, ws.max_row + 1):
        raw = ws.cell(r, 1).value
        if raw is None:
            continue
        match = re.match(r"\s*(\d+)", str(raw))
        if match and int(match.group(1)) == bug_no:
            return r
    return None


def save(wb: Workbook, path: Path = WORKBOOK_PATH) -> Path:
    """
    Save the workbook, falling back to a sibling .tmp file if it got locked
    mid-run.

    Losing a run's results (and the JIRA bug numbers already filed against them)
    to a PermissionError is worse than an awkward filename, so the fallback
    always writes somewhere and reports where.
    """
    path = Path(path)
    try:
        wb.save(path)
        return path
    except PermissionError:
        fallback = path.with_suffix(".tmp.xlsx")
        wb.save(fallback)
        raise WorkbookLockedError(
            f"{path.name} was locked when saving (is it open in Excel?).\n"
            f"Results were written to {fallback} instead — close Excel and merge it in."
        ) from None
